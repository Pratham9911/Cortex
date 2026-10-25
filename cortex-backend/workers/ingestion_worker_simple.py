import sys
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import asyncio
import os
import re
import shutil
import signal
from typing import List, Dict, Any, Tuple

from dotenv import load_dotenv
from bullmq import Worker
from supabase import create_client
from sqlalchemy import text, func

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

from database import SessionLocal
from modelmetrics import sync_project_storage_snapshot
from models import (
    DocumentVersion,
    DocumentChunk,
)
from rag.chunker import chunk_text
from rag.embedder import generate_embeddings

load_dotenv()

REDIS_URL = os.getenv("REDIS_URL")
QUEUE_NAME = "document-ingestion"

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

supabase = create_client(
    SUPABASE_URL,
    SUPABASE_KEY
)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
}


# ============================================================
# LIGHTWEIGHT TEXT EXTRACTION & CHUNKING (NO DOCLING)
# ============================================================

def _extract_pages_from_pdf(file_path: Path) -> List[Tuple[int, str]]:
    """
    Extract text page-by-page from a PDF file using lightweight extractors (PyMuPDF / pypdfium2).
    Preserves 1-based page numbers.
    Explicitly validates against empty files, encryption/passwords, corruption, and scanned PDFs.
    """
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    file_size = file_path.stat().st_size
    if file_size == 0:
        raise ValueError("PDF file is empty (0 bytes).")

    pages_text: List[Tuple[int, str]] = []
    extraction_error = None

    # 1. Primary extractor: PyMuPDF (fitz) - high-speed, C-based, low-memory
    if fitz is not None:
        doc = None
        try:
            doc = fitz.open(str(file_path))
            if doc.is_encrypted:
                if not doc.authenticate(""):
                    raise ValueError("PDF is encrypted or password-protected and cannot be read.")

            if len(doc) == 0:
                raise ValueError("PDF contains no pages.")

            for page_idx, page in enumerate(doc, start=1):
                page_text = page.get_text() or ""
                cleaned = page_text.strip()
                if cleaned:
                    pages_text.append((page_idx, cleaned))

            if pages_text:
                return pages_text
        except ValueError:
            raise
        except Exception as e:
            extraction_error = e
            err_msg = str(e).lower()
            if "password" in err_msg or "encrypted" in err_msg:
                raise ValueError("PDF is encrypted or password-protected and cannot be read.") from e
        finally:
            if doc is not None:
                doc.close()

    # 2. Secondary extractor fallback: pypdfium2 (PDFium engine)
    if not pages_text and pdfium is not None:
        doc = None
        try:
            doc = pdfium.PdfDocument(str(file_path))
            total_pages = len(doc)
            if total_pages == 0:
                raise ValueError("PDF contains no pages.")

            for page_idx in range(total_pages):
                page = doc[page_idx]
                textpage = page.get_textpage()
                page_text = textpage.get_text_bounded() or ""
                cleaned = page_text.strip()
                if cleaned:
                    pages_text.append((page_idx + 1, cleaned))

            if pages_text:
                return pages_text
        except ValueError:
            raise
        except Exception as e:
            err_msg = str(e).lower()
            if "password" in err_msg or "encrypted" in err_msg:
                raise ValueError("PDF is encrypted or password-protected and cannot be read.") from e
            if extraction_error:
                raise ValueError(f"Malformed or corrupted PDF document: {extraction_error} / {e}") from e
            raise ValueError(f"Malformed or corrupted PDF document: {e}") from e
        finally:
            if doc is not None:
                doc.close()

    if extraction_error is not None:
        raise ValueError(f"Malformed or corrupted PDF document: {extraction_error}")

    if not pages_text:
        raise ValueError(
            "No extractable text found in PDF. Scanned or image-only PDFs "
            "without an embedded text layer are not supported by the lightweight text worker."
        )

    return pages_text


def _extract_pages_from_docx(file_path: Path) -> List[Tuple[int, str]]:
    """
    Extract text from a DOCX file using python-docx.
    Extracts paragraph text and table cell text.
    DOCX is continuous flow text without fixed native page numbers,
    so text is assigned page 1.
    """
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    file_size = file_path.stat().st_size
    if file_size == 0:
        raise ValueError("DOCX file is empty (0 bytes).")

    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(file_path))
    except Exception as e:
        raise ValueError(f"Malformed or corrupted DOCX document: {e}") from e

    lines: List[str] = []
    for paragraph in doc.paragraphs:
        txt = paragraph.text.strip()
        if txt:
            lines.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                lines.append(" | ".join(row_texts))

    full_text = "\n\n".join(lines).strip()
    if not full_text:
        raise ValueError("No extractable text found in DOCX document.")

    return [(1, full_text)]


def _extract_pages_from_pptx(file_path: Path) -> List[Tuple[int, str]]:
    """
    Extract text from a PPTX presentation using python-pptx.
    Each slide maps directly to a 1-based page number.
    Extracts text from text frames, shapes, and tables.
    """
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    file_size = file_path.stat().st_size
    if file_size == 0:
        raise ValueError("PPTX file is empty (0 bytes).")

    try:
        from pptx import Presentation
        prs = Presentation(str(file_path))
    except Exception as e:
        raise ValueError(f"Malformed or corrupted PPTX document: {e}") from e

    if not prs.slides:
        raise ValueError("PPTX presentation contains no slides.")

    pages_text: List[Tuple[int, str]] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_lines: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    txt = paragraph.text.strip()
                    if txt:
                        slide_lines.append(txt)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        slide_lines.append(" | ".join(row_texts))

        slide_text = "\n".join(slide_lines).strip()
        if slide_text:
            pages_text.append((slide_idx, slide_text))

    if not pages_text:
        raise ValueError("No extractable text found in PPTX presentation.")

    return pages_text


def _extract_pages_from_text(file_path: Path) -> List[Tuple[int, str]]:
    """
    Extract text from plain text or markdown files.
    All text is assigned page number 1.
    """
    try:
        content = file_path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        raise ValueError(f"Failed to read text document: {str(e)}") from e

    cleaned = content.strip()
    if not cleaned:
        raise ValueError("Text document is empty.")

    return [(1, cleaned)]


def extract_and_chunk_simple(file_path: str | Path) -> List[Dict[str, Any]]:
    """
    Lightweight document extractor and chunker.
    Extracts text page-by-page, adds page markers, applies the existing
    RecursiveCharacterTextSplitter from rag.chunker, and maps chunks to 1-based page numbers.
    Does NOT use Docling or Torch.
    """
    file_path = Path(file_path)
    ext = file_path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: '{ext}'. "
            f"The lightweight worker supports: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    print(f"\n[Lightweight Worker] Extracting: {file_path.name}")

    if ext == ".pdf":
        pages = _extract_pages_from_pdf(file_path)
    elif ext == ".docx":
        pages = _extract_pages_from_docx(file_path)
    elif ext == ".pptx":
        pages = _extract_pages_from_pptx(file_path)
    else:
        pages = _extract_pages_from_text(file_path)

    # Build document text with embedded page markers to track page boundaries through chunker
    document_text = ""
    for page_num, text_content in pages:
        document_text += f"[[PAGE_{page_num}]]\n{text_content}\n"

    # Reuse existing Cortex chunking implementation
    raw_chunks = chunk_text(document_text)

    if not raw_chunks:
        raise ValueError("No chunks were generated from the extracted text.")

    # Map chunks back to 1-based page numbers and clean up page markers
    chunks: List[Dict[str, Any]] = []
    current_page = pages[0][0] if pages else 1

    for raw_chunk in raw_chunks:
        content = raw_chunk["content"]

        page_matches = re.findall(r"\[\[PAGE_(\d+)\]\]", content)
        if page_matches:
            current_page = int(page_matches[0])

        cleaned_content = re.sub(r"\[\[PAGE_\d+\]\]", "", content).strip()

        if cleaned_content:
            chunks.append({
                "chunk_index": len(chunks),
                "content": cleaned_content,
                "page_number": current_page,
            })

    if not chunks:
        raise ValueError("No extractable text found in document after chunking.")

    print(f"[Lightweight Worker] Generated {len(chunks)} chunks across {len(pages)} text pages.")
    return chunks


# ============================================================
# PROCESS DOCUMENT (BULLMQ JOB HANDLER)
# ============================================================

async def process(job, job_token):
    print("\n==========================================")
    print("LIGHTWEIGHT INGESTION JOB STARTED")
    print("==========================================")
    print("Job ID:", job.id)
    print("Job Name:", job.name)
    print("Job Data:", job.data)

    document_id = job.data["document_id"]
    version_id = job.data["version_id"]
    project_id = job.data["project_id"]
    storage_path = job.data["storage_path"]

    db = SessionLocal()

    temp_dir = (
        Path("worker_temp")
        / f"document_{document_id}"
        / f"version_{version_id}"
    )

    local_path = None

    try:
        # ==================================================
        # 1. MARK PROCESSING
        # ==================================================
        version = db.query(DocumentVersion).filter(
            DocumentVersion.version_id == version_id
        ).first()

        if not version:
            raise RuntimeError(f"DocumentVersion {version_id} not found")

        if version.status == "completed":
            print(f"Version {version_id} already marked completed. Skipping redundant processing.")
            return {
                "status": "completed",
                "document_id": document_id,
                "version_id": version_id,
                "already_processed": True,
            }

        version.status = "processing"
        db.commit()

        # ==================================================
        # 2. DOWNLOAD FROM SUPABASE STORAGE
        # ==================================================
        print("\n========== DOWNLOAD ==========")
        print("Storage path:", storage_path)

        file_bytes = (
            supabase.storage
            .from_("documents")
            .download(storage_path)
        )

        print(f"Downloaded: {len(file_bytes)} bytes")

        # ==================================================
        # 3. WRITE TO LOCAL FILE
        # ==================================================
        temp_dir.mkdir(parents=True, exist_ok=True)
        file_name = Path(storage_path).name
        local_path = temp_dir / file_name
        local_path.write_bytes(file_bytes)

        print("Local file:", local_path.resolve())

        # ==================================================
        # 4. LIGHTWEIGHT EXTRACTION & CHUNKING
        # ==================================================
        print("\n========== EXTRACTION & CHUNKING ==========")
        chunks = await asyncio.to_thread(
            extract_and_chunk_simple,
            local_path
        )

        if not chunks:
            raise RuntimeError("No chunks were generated")

        print(f"Chunks generated: {len(chunks)}")

        # ==================================================
        # 5. EMBEDDINGS
        # ==================================================
        print("\n========== EMBEDDING ==========")
        embedded_chunks = await asyncio.to_thread(
            generate_embeddings,
            chunks,
            20
        )

        if not embedded_chunks:
            raise RuntimeError("No embedded chunks were generated")

        print(f"Embedded chunks: {len(embedded_chunks)}")

        # ==================================================
        # 6. DATABASE INSERT (IDEMPOTENT)
        # ==================================================
        print("\n========== DATABASE INSERT ==========")

        # Clean up any partial/pre-existing chunks for this version to ensure
        # retried attempts are strictly idempotent and never duplicate data
        deleted_count = db.query(DocumentChunk).filter(
            DocumentChunk.version_id == version_id
        ).delete()

        if deleted_count > 0:
            print(f"Cleared {deleted_count} stale chunks for version {version_id} before re-inserting.")
            db.flush()

        for chunk in embedded_chunks:
            db_chunk = DocumentChunk(
                project_id=project_id,
                version_id=version_id,
                chunk_index=chunk["chunk_index"],
                content=chunk["content"],
                page_number=chunk.get("page_number"),
                embedding=chunk["embedding"]
            )
            db.add(db_chunk)

        db.flush()

        # ==================================================
        # 7. UPDATE SEARCH VECTOR
        # ==================================================
        db.execute(
            text("""
                UPDATE document_chunks
                SET search_vector =
                    to_tsvector(
                        'english',
                        content
                    )
                WHERE version_id = :version_id
            """),
            {
                "version_id": version_id
            }
        )

        # ==================================================
        # 8. MARK VERSION COMPLETED
        # ==================================================
        version.status = "completed"
        db.flush()

        # ==================================================
        # 9. ACTIVATE LATEST COMPLETED VERSION
        # ==================================================
        latest_completed_version = db.query(
            DocumentVersion
        ).filter(
            DocumentVersion.document_id == document_id,
            DocumentVersion.is_deleted == False,
            DocumentVersion.status == "completed"
        ).order_by(
            DocumentVersion.version_number.desc()
        ).first()

        if not latest_completed_version:
            raise RuntimeError("No completed version found")

        # Deactivate all other active versions
        active_versions = db.query(
            DocumentVersion
        ).filter(
            DocumentVersion.document_id == document_id,
            DocumentVersion.is_active == True,
            DocumentVersion.is_deleted == False,
            DocumentVersion.version_id != latest_completed_version.version_id
        ).all()

        for active_version in active_versions:
            active_version.is_active = False

        # Activate the highest completed version
        latest_completed_version.is_active = True
        latest_completed_version.activated_by = latest_completed_version.uploaded_by
        latest_completed_version.activated_at = func.now()

        sync_project_storage_snapshot(db, project_id)
        db.flush()

        # ==================================================
        # 10. COMMIT TRANSACTION
        # ==================================================
        db.commit()

        print("\n==========================================")
        print("LIGHTWEIGHT INGESTION COMPLETED")
        print("==========================================")
        print("Document:", document_id)
        print("Version:", version_id)
        print("Chunks:", len(embedded_chunks))
        print("Active version:", latest_completed_version.version_number)

        return {
            "status": "completed",
            "document_id": document_id,
            "version_id": version_id,
            "chunks": len(embedded_chunks),
            "active_version": latest_completed_version.version_number
        }

    # ======================================================
    # FAILURE HANDLING
    # ======================================================
    except Exception as e:
        print("\n==========================================")
        print("LIGHTWEIGHT INGESTION FAILED")
        print("==========================================")
        print("Document:", document_id)
        print("Version:", version_id)
        print("Error:", str(e))

        db.rollback()

        # Record failure state in DB so the UI/frontend accurately reflects failed state
        try:
            version = db.query(
                DocumentVersion
            ).filter(
                DocumentVersion.version_id == version_id
            ).first()

            if version:
                version.status = "failed"
                # A failed version must never become or remain active
                version.is_active = False
                db.commit()
                print(f"Version {version_id} -> marked as failed")

        except Exception as status_error:
            db.rollback()
            print("Failed to update document status to failed:", status_error)

        # Re-raise exception so BullMQ acknowledges the attempt failure and schedules retries
        raise

    finally:
        # ==================================================
        # 11. CLOSE DB
        # ==================================================
        db.close()

        # ==================================================
        # 12. CLEAN UP TEMPORARY DIRECTORY
        # ==================================================
        if temp_dir.exists():
            try:
                shutil.rmtree(temp_dir, ignore_errors=True)
                print(f"Temporary version directory cleaned: {temp_dir}")
            except Exception as cleanup_error:
                print("Temporary cleanup failed:", cleanup_error)


# ============================================================
# WORKER LIFECYCLE
# ============================================================

async def main():
    shutdown_event = asyncio.Event()

    def signal_handler(sig, frame):
        print(f"\n[Lightweight Worker] Shutdown signal ({sig}) received.")
        shutdown_event.set()

    signal.signal(signal.SIGINT, signal_handler)
    signal.signal(signal.SIGTERM, signal_handler)

    concurrency = int(os.getenv("WORKER_CONCURRENCY", "1"))

    worker = Worker(
        QUEUE_NAME,
        process,
        {
            "connection": REDIS_URL,
            "concurrency": concurrency,
        }
    )

    print(f"[Lightweight Worker] Listening on queue: '{QUEUE_NAME}' (concurrency={concurrency})")

    await shutdown_event.wait()

    print("[Lightweight Worker] Closing worker gracefully...")
    await worker.close()
    print("[Lightweight Worker] Worker stopped cleanly.")


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    asyncio.run(main())
